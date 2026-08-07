import os
import sys
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Border, Side, PatternFill
from openpyxl.utils import get_column_letter


def get_app_dir():
    if getattr(sys, "frozen", False):
        return os.path.dirname(os.path.abspath(sys.executable))
    return os.path.dirname(os.path.abspath(__file__))


app_dir = get_app_dir()


def find_input_workbook():
    today = datetime.now().strftime('%d-%m-%y')
    candidates = []

    for base in [Path.cwd(), Path(app_dir), Path(__file__).resolve().parent]:
        candidates.append(base / f"Carteira_Fictícia_{today}.xlsx")
        candidates.append(base / "Carteira_Fictícia.xlsx")

    for pattern in ["Carteira_Fictícia*.xlsx", "Carteira*.xlsx"]:
        for base in [Path.cwd(), Path(app_dir), Path(__file__).resolve().parent]:
            candidates.extend(base.glob(pattern))

    seen = set()
    results = []
    for path in candidates:
        if not path.exists():
            continue
        if path in seen:
            continue
        seen.add(path)
        results.append(path)

    if not results:
        return None

    return max(results, key=lambda p: p.stat().st_mtime)


def ensure_input_workbook():
    workbook = find_input_workbook()
    if workbook is not None:
        return str(workbook)

    try:
        from main import main as gerar_carteira
        gerar_carteira()
    except Exception:
        pass

    workbook = find_input_workbook()
    if workbook is not None:
        return str(workbook)

    raise FileNotFoundError("Arquivo Carteira_Fictícia*.xlsx não encontrado. Gere a carteira primeiro.")


def mainFup():
    hoje = datetime.now()
    semana_atual = hoje - timedelta(days=hoje.weekday())  # Segunda-feira da semana atual
    semana_passada = semana_atual - timedelta(days=7)  # Segunda-feira da

    input_workbook = ensure_input_workbook()
    print(f"Usando arquivo de entrada: {input_workbook}")

    # Carregar o workbook original
    wb = load_workbook(input_workbook)

    # Renomear planilha se necessário
    output_workbook = os.path.join(app_dir, "Carteira_Test.xlsx")
    if 'Export' in wb.sheetnames:
        ws = wb['Export']
        ws.title = "PANELA"
        wb.save(output_workbook)

    wb.save(output_workbook)

    # Função para processar uma aba
    def processar_aba(aba_nome, pasta):
        # Importar o arquivo Excel (mantendo todas as colunas, incluindo FABRICANTE)
        df = pd.read_excel(output_workbook, sheet_name=aba_nome)

        # Excluir linhas com STATUS PEDIDO igual a Concluído
        if 'Status Pedido' in df.columns:
            df = df[df['Status Pedido'] != 'Concluído']

        if 'OBS' in df.columns:
            df = df[df['OBS'] != 'Enviado']

        if 'Data Emissão' in df.columns:
            df['Data Emissão'] = pd.to_datetime(df['Data Emissão'], errors='coerce')
            df = df[df['Data Emissão'] <= semana_passada]

        # Colunas necessárias
        colunas_necessarias = ['Griffe', 'Cliente', 'Pedido ID', 'Data Original Entrega', 'Data Emissão', 'Data Confirmação', 'Referência', 'Descrição', 'Linha']

        # Para CASTRO, adicionar as colunas extras
        if aba_nome.upper() == 'CASTRO':
            colunas_necessarias += ['MOTIVO DA PRORROGAÇÃO', 'OBS GERAIS']
            for col in ['MOTIVO DA PRORROGAÇÃO', 'OBS GERAIS']:
                if col not in df.columns:
                    df[col] = ''

        df = df[colunas_necessarias]

        df['STATUS DE PRODUÇÃO'] = ''
        df['DATA DE ENTREGA'] = ''

        def salvar_e_estilizar(grupo_df, nome_arquivo):
            if grupo_df.empty:
                print(f"Pulando exportação sem dados: {nome_arquivo}")
                return

            if 'Cliente' in grupo_df.columns:
                grupo_df = grupo_df.drop(columns=['Cliente'])

            os.makedirs(os.path.dirname(nome_arquivo), exist_ok=True)
            with pd.ExcelWriter(nome_arquivo, engine='openpyxl') as writer:
                grupo_df.to_excel(writer, sheet_name='Sheet1', index=False)

            wb_novo = load_workbook(nome_arquivo)
            ws = wb_novo.active

            if ws.max_row > 1 and ws.max_column > 0:
                first_col = 1
                last_col = ws.max_column
                ws.auto_filter.ref = f"A1:{get_column_letter(last_col)}1"

            dims = {}
            for row in ws.rows:
                for cell in row:
                    if cell.value:
                        dims[cell.column] = max((dims.get(cell.column, 0), len(str(cell.value))))
            for col, value in dims.items():
                ws.column_dimensions[get_column_letter(col)].width = value + 2

            borda = Border(left=Side(border_style='thin'), right=Side(border_style='thin'), top=Side(border_style='thin'), bottom=Side(border_style='thin'))
            for linha in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
                for celula in linha:
                    celula.border = borda

            fill = PatternFill(start_color='FABF8F', end_color='FABF8F', fill_type='solid')
            header_row = 1
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=header_row, column=col)
                if cell.value is not None:
                    cell.fill = fill

            # Formatar PEDIDO
            pedido_col = None
            for col in range(1, ws.max_column + 1):
                if ws.cell(row=1, column=col).value == 'Pedido ID':
                    pedido_col = col
                    break
            if pedido_col:
                for row in range(2, ws.max_row + 1):
                    cell = ws.cell(row=row, column=pedido_col)
                    if cell.value is not None:
                        cell.number_format = '00000000'

            # Formatar datas
            colunas_datas = ['Data Original Entrega', 'Data Emissão', 'Data Confirmação']
            formato_data = 'DD/MM/YYYY'
            for nome_coluna in colunas_datas:
                data_col = None
                for col in range(1, ws.max_column + 1):
                    if ws.cell(row=1, column=col).value == nome_coluna:
                        data_col = col
                        break
                if data_col:
                    for row in range(2, ws.max_row + 1):
                        cell = ws.cell(row=row, column=data_col)
                        if cell.value is not None:
                            cell.number_format = formato_data

            wb_novo.save(nome_arquivo)

        # Agrupar por Cliente
        for cliente, grupo in df.groupby('Cliente'):
            # Verificar griffes no grupo e criar abreviações
            griffes_presentes = set(grupo['Griffe'].unique())
            abreviacoes = []
            if any('Aura & Co' in g for g in griffes_presentes):
                abreviacoes.append('AC')
            if any('L\'Éclat' in g for g in griffes_presentes):
                abreviacoes.append('EC')
            if any('Vanguardia' in g for g in griffes_presentes):
                abreviacoes.append('VD')

            sufixo = '' + ' & '.join(abreviacoes) if abreviacoes else ''

            if cliente.strip().lower() == 'epos jeans':
                filial_text = grupo['FILIAL'].astype(str)
                grupo_a = grupo[filial_text.str.contains('Mostru[aá]rio|Desfile', case=False, na=False)]
                grupo_b = grupo[filial_text.str.contains('Recebimento', case=False, na=False)]

                if not grupo_a.empty:
                    nome_arquivo = f"{pasta}/{sufixo} - FUP - {cliente} - MOSTRUARIO_DESFILE - {datetime.now().strftime('%d.%m')}.xlsx"
                    salvar_e_estilizar(grupo_a, nome_arquivo)

                if not grupo_b.empty:
                    nome_arquivo = f"{pasta}/{sufixo} - FUP - {cliente} - PRODUÇÃO - {datetime.now().strftime('%d.%m')}.xlsx"
                    salvar_e_estilizar(grupo_b, nome_arquivo)

                # Também garante que o grupo total não é salvo novamente para evitar duplicação.
                continue

            nome_arquivo = f"{pasta}/{sufixo} - FUP - {cliente} - {datetime.now().strftime('%d.%m')}.xlsx"
            salvar_e_estilizar(grupo, nome_arquivo)



    os.makedirs(os.path.join(app_dir, 'follow ups panela'), exist_ok=True)
    os.makedirs(os.path.join(app_dir, 'follow ups castro'), exist_ok=True)

    # Processar PANELA
    processar_aba('PANELA', os.path.join(app_dir, 'follow ups panela'))

    # Processar CASTRO
    processar_aba('CASTRO', os.path.join(app_dir, 'follow ups castro'))

    print("Planilhas criadas para PANELA e CASTRO.")


if __name__ == "__main__":
    mainFup()